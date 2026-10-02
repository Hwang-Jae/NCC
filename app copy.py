import os
import streamlit as st
import clickhouse_connect
import ollama
from webdav3.client import Client as WebDavClient
import pdfplumber
from docx import Document
import openpyxl

# --- 1. 설정 정보 ---
CLOUDDOC_OPTIONS = {
    'webdav_url': 'http://idisk.nvhkorea.com:80',
    'webdav_username': 'jwhwang1',
    'webdav_password': 'jwhwang1'
}
EMBED_MODEL = 'nomic-embed-text'  # 벡터 변환용 임베딩 모델 (Ollama에서 미리 pull 받아주세요)
LLM_MODEL = 'gpt-oss:20b'        # 답변 생성용 LLM

# --- 2. ClickHouse 연결 설정 ---
def get_db_client():
    return clickhouse_connect.get_client(
        host='localhost', port=8123, username='default', password='clickhouse',
        connect_timeout=30, send_receive_timeout=300
    )

# DB 테이블 초기화 (벡터 검색용 Array(Float32) 포함)
def init_clickhouse():
    client = get_db_client()
    client.command("""
        CREATE TABLE IF NOT EXISTS doc_chunks (
            id String,
            file_name String,
            content String,
            embedding Array(Float32)
        ) ENGINE = MergeTree() ORDER BY id
    """)

# --- 3. 파일 텍스트 추출 함수들 ---
def extract_text(file_path):
    ext = os.path.splitext(file_path)[-1].lower()
    text = ""
    try:
        if ext == '.pdf':
            with pdfplumber.open(file_path) as pdf:
                text = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
        elif ext == '.docx':
            doc = Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs])
        elif ext in ['.xlsx', '.xls']:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.sheetnames:
                for row in wb[sheet].iter_rows(values_only=True):
                    text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
        elif ext == '.txt':
            with open(file_path, 'r', encoding='utf-8') as f:
                text = f.read()
    except Exception as e:
        print(f"파일 읽기 실패 ({file_path}): {e}")
    return text

# 텍스트 일정한 크기로 자르기 (Chunking)
def chunk_text(text, chunk_size=500):
    return [text[i:i+chunk_size] for i in range(0, len(text), chunk_size)]

# --- 4. 클라우드독에서 파일 다운로드 및 DB 동기화 ---
def sync_cloud_doc():
    st.info("클라우드독에서 파일을 가져오는 중입니다...")
    try:
        webdav_client = WebDavClient(CLOUDDOC_OPTIONS)
        # 클라우드독 루트 폴더의 파일 목록 가져오기 (필요시 하위 경로 지정)
        files = webdav_client.list()
    except Exception as e:
        st.error(f"클라우드독 연결 실패: {e}")
        return

    client = get_db_client()
    # 기존 데이터 초기화 (재동기화 시 시나리오)
    client.command("TRUNCATE TABLE doc_chunks")

    os.makedirs("./temp_docs", exist_ok=True)
    
    for file_name in files:
        if file_name.endswith(('/', '.', '..')): continue # 디렉토리 제외
        
        local_path = os.path.join("./temp_docs", file_name)
        webdav_client.download_file(file_name, local_path)
        
        # 텍스트 추출 및 분할
        full_text = extract_text(local_path)
        chunks = chunk_text(full_text)
        
        if not chunks: continue
        
        # 벡터 변환 및 ClickHouse 적재
        data_to_insert = []
        for i, chunk in enumerate(chunks):
            # Ollama를 이용한 임베딩 생성
            response = ollama.embed(model=EMBED_MODEL, input=chunk)
            embedding = response['embeddings'][0]
            
            chunk_id = f"{file_name}_{i}"
            data_to_insert.append([chunk_id, file_name, chunk, embedding])
        
        if data_to_insert:
            client.insert('doc_chunks', data_to_insert, column_names=['id', 'file_name', 'content', 'embedding'])
            
    st.success("모든 문서가 ClickHouse에 고속 저장되었습니다!")

# --- 5. RAG 검색 및 Ollama 답변 생성 ---
def ask_question(query):
    # 1. 질문 벡터화
    query_embed = ollama.embed(model=EMBED_MODEL, input=query)['embeddings'][0]
    
    # 2. ClickHouse 코사인 유사도 고속 검색 (Standard Array distance 함수 활용)
    client = get_db_client()
    search_query = """
        SELECT content, file_name, cosineDistance(embedding, {q}) as distance 
        FROM doc_chunks 
        ORDER BY distance ASC 
        LIMIT 3
    """
    result = client.query(search_query, parameters={'q': query_embed})
    
    # 3. 검색된 참고 문서 컨텍스트 생성
    context = ""
    sources = set()
    for row in result.result_rows:
        context += f"[{row[1]} 내용]: {row[0]}\n\n"
        sources.add(row[1])
        
    if not context:
        return "참고할 만한 문서 내용을 찾지 못했습니다.", []

    # 4. Ollama gpt-oss:20b 로 답변 생성
    system_prompt = f"""당신은 사내 문서 정보를 바탕으로 답변하는 유능한 어시스턴트입니다. 
주어진 [참고 문서] 내용에 기반하여 사용자의 질문에 정확하고 친절하게 답하세요. 문서에 없는 내용은 억지로 지어내지 마세요.

[참고 문서]
{context}
"""
    
    response = ollama.generate(model=LLM_MODEL, prompt=f"{system_prompt}\n\n질문: {query}")
    return response['response'], list(sources)

# --- 6. Streamlit 웹 UI 구성 ---
st.set_page_config(page_title="사내 문서 Q&A 시스템", layout="wide")
init_clickhouse()

st.title("📁 사내 공유디스크 지식인 AI")
st.subheader("ClickHouse + Ollama 고속 문서 검색 시스템")

# 사이드바 - 동기화 버튼
with st.sidebar:
    st.header("⚙️ 데이터 관리")
    if st.button("🔄 클라우드독 파일 동기화시작", use_container_width=True):
        with st.spinner("문서 분석 중..."):
            sync_cloud_doc()
    st.caption("공유 디스크의 도면, 인프라 구성도, IP 대역 표 등 최신 파일을 가져와 DB를 업데이트합니다.")

# 메인 화면 - 채팅 인터페이스
if "messages" not in st.session_state:
    st.session_state.messages = []

# 기존 대화 출력
for msg in st.session_state.messages:
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])

# 사용자 입력
if user_query := st.chat_input("ex) VW 인프라 구성도를 보고싶어요 / 경주공장 사무실 대역 IP 알려줘"):
    st.session_state.messages.append({"role": "user", "content": user_query})
    with st.chat_message("user"):
        st.markdown(user_query)
        
    with st.chat_message("assistant"):
        with st.spinner("ClickHouse 검색 및 답변 생성 중..."):
            answer, source_files = ask_question(user_query)
            
            if source_files:
                answer += f"\n\n--- \n📌 **참고한 파일:** {', '.join(source_files)}"
                
            st.markdown(answer)
            st.session_state.messages.append({"role": "assistant", "content": answer})